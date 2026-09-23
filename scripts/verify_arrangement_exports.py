"""Read-only independent QA of browser downloads; run with openpyxl installed.
Usage: python scripts/verify_arrangement_exports.py frontend/test-results/grid-export-final
"""
import json
import sys
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COLORS = ['FFFFFF', 'EDEDED', 'D9D9D9', 'C4C4C4']

def verify(path):
    snapshot = json.loads(path.with_suffix('.json').read_text(encoding='utf-8'))
    wb = load_workbook(path)
    assert wb.sheetnames == ['比賽安排']
    ws = wb.active
    assert not ws.protection.sheet
    assert ws['A1'].value == snapshot['title'] and ws['A2'].value == snapshot['version']
    people = {r['registration_id']: r for r in snapshot['rows']}
    layout = snapshot['layout']
    if layout is None:
        assert '不代表當時分組' in ws['A3'].value
        groups = [sorted([r for r in people.values() if r['competition_level'] == level], key=lambda r:r['queue_sequence']) for level in range(1,11)]
        layout = dict(columns=[dict(id=str(i),kind='level',level=i+1) for i in range(10)], rows=[dict(id=str(i),role='body') for i in range(max(8,max(map(len,groups))))], merges=[], cells=[dict(row_id=str(j),column_id=str(i),kind='registration',registration_id=r['registration_id']) for i,g in enumerate(groups) for j,r in enumerate(g)])
    rows, cols = layout['rows'], layout['columns']
    source = {(c['row_id'],c['column_id']):c for c in layout['cells']}
    local = {(c['row_id'],c['column_id']):c['shade'] for c in layout.get('cell_shades',[])}
    expected, fills, mapped = {}, {}, {}
    def put(r,c,value,shade):
        expected[r,c]=value; fills[r,c]=COLORS[shade]
    header_rows=[]
    cursor=4
    for ri,row in enumerate(rows):
        if row['role']=='body' and (ri==0 or rows[ri-1]['role']=='header'):
            header_rows.append(cursor)
            for ci,col in enumerate(cols,1):
                title=col.get('title')
                if title is None: title=str(col['level'])+' 級' if col['kind']=='level' else '文字／備註'
                put(cursor,ci,title,col['header_shade'] if col.get('header_shade') is not None else col.get('shade',0))
            cursor+=1
        mapped[row['id']]=cursor
        for ci,col in enumerate(cols,1):
            point=(row['id'],col['id']); cell=source.get(point,{})
            value=cell.get('text','')
            if cell.get('kind')=='registration':
                person=people[cell['registration_id']]
                value=person['member_name']+(' [素]' if person['diet']=='vegetarian' else '')+('\n'+person['distinguishing_note'] if person.get('distinguishing_note') else '')
            put(cursor,ci,value,local.get(point,max(row.get('shade',0),col.get('shade',0))))
        cursor+=1
    expected_merges={f'A{i}:{get_column_letter(len(cols))}{i}' for i in range(1,4)}
    row_ids=[r['id'] for r in rows]; col_ids=[c['id'] for c in cols]
    for header_row in header_rows:
        for merge in layout.get('header_merges',[]):
            left=col_ids.index(merge['start_column_id'])+1;right=col_ids.index(merge['end_column_id'])+1
            coords=[(header_row,c) for c in range(left,right+1)]
            title=merge.get('title')
            if title is None:title=' '.join(expected[p] for p in coords)
            fill=max((fills[p] for p in coords),key=COLORS.index)
            for p in coords:expected.pop(p);fills.pop(p)
            put(header_row,left,title,COLORS.index(fill))
            expected_merges.add(f'{get_column_letter(left)}{header_row}:{get_column_letter(right)}{header_row}')
    for merge in layout['merges']:
        top,bottom=[row_ids.index(merge[k]['row_id']) for k in ('start','end')]
        left,right=[col_ids.index(merge[k]['column_id'])+1 for k in ('start','end')]
        first,last=mapped[rows[top]['id']],mapped[rows[bottom]['id']]
        coords=[(mapped[rows[r]['id']],c) for r in range(top,bottom+1) for c in range(left,right+1)]
        text=' '.join(expected[p] for p in coords if expected[p])
        fill=max((fills[p] for p in coords),key=COLORS.index)
        for p in coords: expected.pop(p); fills.pop(p)
        put(first,left,text,COLORS.index(fill))
        expected_merges.add(f'{get_column_letter(left)}{first}:{get_column_letter(right)}{last}')
    assert set(map(str,ws.merged_cells.ranges))==expected_merges
    assert (ws.max_row,ws.max_column)==(cursor-1,len(cols))
    for (r,c),value in expected.items():
        cell=ws.cell(r,c)
        assert (cell.value or '')==value,(path.name,cell.coordinate,cell.value,value)
        assert cell.fill.fgColor.rgb[-6:]==fills[r,c],(cell.coordinate,cell.fill.fgColor)
        assert cell.alignment.wrap_text and cell.number_format=='@'
        assert all(getattr(cell.border,side).style=='thin' for side in ('top','bottom','left','right'))
        assert ws.row_dimensions[r].height>=30
    for row in ws:
        for cell in row:
            assert cell.data_type!='f' and not cell.hyperlink
            if cell.value is not None: assert cell.data_type=='s'
    assert all(ws.column_dimensions[get_column_letter(c)].width==20 for c in range(1,len(cols)+1))
    values=[c.value for row in ws for c in row if isinstance(c.value,str)]
    for person in people.values():
        assert sum(v.split('\n')[0].removesuffix(' [素]')==person['member_name'] for v in values)==1
    with ZipFile(path) as z:
        assert not any(any(tag in n.lower() for tag in ('vbaproject','externallinks','xl/media/')) for n in z.namelist())
        assert not any(b'TargetMode="External"' in z.read(n) for n in z.namelist() if n.endswith('.rels'))
    return dict(file=str(path),people=len(people),rows=ws.max_row,columns=ws.max_column,merges=len(expected_merges),cells_checked=len(expected),status='PASS')

if __name__=='__main__':
    paths=sorted(Path(sys.argv[1]).rglob('*.xlsx'))
    assert paths, 'No downloaded workbooks found'
    print(json.dumps([verify(p) for p in paths],ensure_ascii=False,indent=2))
